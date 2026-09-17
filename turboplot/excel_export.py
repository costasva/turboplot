"""Export a plot to Excel: the data it was built from, plus the plot itself
as a native Excel chart the recipient can edit and re-style.

Layout of the workbook:
  Chart  the chart, sized to the page
  Data   the numbers behind it, one column pair (x, y) per series
  Info   what the plot is, what was selected, and where each series came from
"""
from __future__ import annotations

import datetime as _dt

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series as ChartSeries
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import PlotSpec
from .plotmodel import PlotData
from .state import SessionState

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="EFEFEC")
PT = 12700                      # EMU per point, for line widths


def _num_format(values: list[float]) -> str:
    """Enough decimals to be useful without turning into noise."""
    biggest = max((abs(v) for v in values if v is not None), default=1.0)
    if biggest >= 100:
        return "0.00"
    if biggest >= 1:
        return "0.0000"
    return "0.000000"


def _line(color: str, dash: str, width_pt: float = 1.75) -> GraphicalProperties:
    return GraphicalProperties(ln=LineProperties(solidFill=color, prstDash=dash,
                                                 w=int(width_pt * PT)))


def _marker(symbol: str, color: str, size: int = 6, hollow: bool = False) -> Marker:
    if symbol in (None, "", "none"):
        return Marker(symbol="none")
    fill = "FFFFFF" if hollow else color
    return Marker(symbol=symbol, size=size,
                  spPr=GraphicalProperties(solidFill=fill,
                                           ln=LineProperties(solidFill=color, w=int(1.1 * PT))))


def _write_header(ws, row: int, col: int, text: str, width: float | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(name=FONT, bold=True, size=10)
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width


def write_plot_workbook(path, state: SessionState, spec: PlotSpec, data: PlotData,
                        chart_title: str, ranges: dict | None = None) -> None:
    wb = Workbook()
    ws_chart = wb.active
    ws_chart.title = "Chart"
    ws_data = wb.create_sheet("Data")
    ws_info = wb.create_sheet("Info")

    if data.kind == "scalar":
        chart = _scalar_sheet(ws_data, state, spec, data)
    else:
        chart = _xy_sheet(ws_data, state, spec, data)

    chart.title = chart_title
    chart.height, chart.width = 14.0, 26.0
    chart.legend.position = "r"
    chart.x_axis.delete = False        # openpyxl hides axes by default in some viewers
    chart.y_axis.delete = False
    # Match the app's axis ranges; without this Excel anchors value axes at zero.
    for name, axis in (("x", chart.x_axis), ("y", chart.y_axis)):
        if name == "x" and data.kind == "scalar":
            continue                   # design IDs are categories, not a scale
        bounds = (ranges or {}).get(name)
        if not bounds:
            continue
        low, high, unit = bounds
        axis.scaling.min, axis.scaling.max = low, high
        if unit:
            axis.majorUnit = unit
    ws_chart.add_chart(chart, "B2")
    ws_chart.sheet_view.showGridLines = False

    _info_sheet(ws_info, state, spec, data, chart_title)
    wb.save(str(path))


# --------------------------------------------------------------------------
# scalar quantity against design ID: a category chart, references as levels
# --------------------------------------------------------------------------
def _scalar_sheet(ws, state: SessionState, spec: PlotSpec, data: PlotData) -> LineChart:
    _write_header(ws, 1, 1, spec.x_label, width=18)
    _write_header(ws, 1, 2, spec.y_label, width=22)
    for j, ref in enumerate(data.refs):
        _write_header(ws, 1, 3 + j, ref.label, width=26)

    values = [s.value for s in data.series]
    fmt = _num_format(values + [r.value for r in data.refs])
    for i, entry in enumerate(data.series, start=2):
        name = ws.cell(row=i, column=1, value=entry.name)
        name.font = Font(name=FONT, size=10)
        cell = ws.cell(row=i, column=2, value=entry.value)
        cell.font, cell.number_format = Font(name=FONT, size=10), fmt
        # A reference is one number; repeating it down the column draws it as
        # the horizontal line the on-screen plot shows.
        for j, ref in enumerate(data.refs):
            rc = ws.cell(row=i, column=3 + j, value=ref.value)
            rc.font, rc.number_format = Font(name=FONT, size=10), fmt
    last = len(data.series) + 1
    ws.freeze_panes = "A2"

    chart = LineChart()
    chart.y_axis.title = spec.y_label
    chart.x_axis.title = spec.x_label

    # The column header names the quantity; the legend entry names the series.
    ser = ChartSeries(Reference(ws, min_col=2, min_row=2, max_row=last),
                      title="Design iterations")
    ser.graphicalProperties = _line(state.theme.c["axis"].lstrip("#").upper(), "solid", 1.25)
    ser.marker = Marker(symbol="none")
    ser.smooth = False
    # One coloured marker per design iteration, matching the plot's colours.
    ser.data_points = [
        DataPoint(idx=i, marker=_marker("circle", state.style_for(e.source).rgb, size=8))
        for i, e in enumerate(data.series)
    ]
    chart.series.append(ser)

    for j, ref in enumerate(data.refs):
        st = state.style_for(ref.source)
        rs = ChartSeries(Reference(ws, min_col=3 + j, min_row=1, max_row=last),
                         title_from_data=True)
        rs.graphicalProperties = _line(st.rgb, st.dash)
        rs.marker = Marker(symbol="none")
        rs.smooth = False
        chart.series.append(rs)

    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    return chart


# --------------------------------------------------------------------------
# curves and radial profiles: an XY chart, one column pair per series
# --------------------------------------------------------------------------
def _xy_sheet(ws, state: SessionState, spec: PlotSpec, data: PlotData) -> ScatterChart:
    profile = data.kind == "profile"
    # Radial results are read with span on the vertical axis, so the quantity
    # is the x column there and the span fraction is y.
    x_label = spec.y_label if profile else spec.x_label
    y_label = spec.x_label if profile else spec.y_label

    chart = ScatterChart()
    chart.x_axis.title = x_label
    chart.y_axis.title = y_label
    if profile:
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 1
        chart.y_axis.majorUnit = 0.25

    multi_op = len(data.ops) > 1
    col = 1
    for entry in data.series:
        st = state.style_for(entry.source)
        _write_header(ws, 1, col, entry.label, width=18)
        _write_header(ws, 1, col + 1, "", width=18)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        _write_header(ws, 2, col, x_label)
        _write_header(ws, 2, col + 1, y_label)
        fmt_x, fmt_y = _num_format(entry.x), _num_format(entry.y)
        for i, (xv, yv) in enumerate(zip(entry.x, entry.y), start=3):
            cx = ws.cell(row=i, column=col, value=xv)
            cy = ws.cell(row=i, column=col + 1, value=yv)
            cx.font = cy.font = Font(name=FONT, size=10)
            cx.number_format, cy.number_format = fmt_x, fmt_y
        last = len(entry.x) + 2

        ser = ChartSeries(Reference(ws, min_col=col + 1, min_row=3, max_row=last),
                          Reference(ws, min_col=col, min_row=3, max_row=last),
                          title=entry.label)
        dash = state.theme.op_dash(data.ops.index(entry.op)) if multi_op else st.dash
        ser.graphicalProperties = _line(st.rgb, dash, 2.0 if entry.source.is_design else 1.75)
        # Reference data keeps its open markers; design profiles are drawn clean.
        show_marker = not (profile and entry.source.is_design)
        ser.marker = _marker(st.symbol if show_marker else "none", st.rgb,
                             size=6, hollow=not entry.source.is_design)
        ser.smooth = False
        chart.series.append(ser)
        col += 2

    ws.freeze_panes = "A3"
    return chart


# --------------------------------------------------------------------------
def _info_sheet(ws, state: SessionState, spec: PlotSpec, data: PlotData,
                chart_title: str) -> None:
    ws.column_dimensions["A"].width = 22
    for letter, width in zip("BCDEF", (34, 16, 26, 14, 60)):
        ws.column_dimensions[letter].width = width
    ws.sheet_view.showGridLines = False

    rows = [
        ("Plot", chart_title),
        ("Component", spec.comp_name),
        ("Quantity", spec.title),
        ("Horizontal axis", spec.y_label if data.kind == "profile" else spec.x_label),
        ("Vertical axis", spec.x_label if data.kind == "profile" else spec.y_label),
        ("Operating points", ", ".join(data.ops) if data.ops else "n/a"),
        ("Exported", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Exported by", "TurboPlot (mockup)"),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        key = ws.cell(row=r, column=1, value=label)
        key.font = Font(name=FONT, bold=True, size=10)
        ws.cell(row=r, column=2, value=value).font = Font(name=FONT, size=10)

    start = len(rows) + 2
    for c, head in enumerate(("Series", "Type", "Hardware", "Date", "Colour", "Notes"), start=1):
        _write_header(ws, start, c, head)
    seen, r = set(), start + 1
    for entry in list(data.series) + list(data.refs):
        src = entry.source
        if src.source_id in seen:
            continue
        seen.add(src.source_id)
        kind = "Design iteration" if src.is_design else f"Reference — {src.ref_type}"
        for c, value in enumerate((src.label, kind, src.hardware, src.created,
                                   "#" + state.style_for(src).rgb, src.notes), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 6))
        r += 1
